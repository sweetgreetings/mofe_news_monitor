# 「📋 엑셀파일로 보고용 텍스트 생성」(`GET /excel-merge`)의 로직 — 화면은 app/excel_merge_renderer.py.
#
# 담당자가 들고 있는 표(언론사·제목·링크)를 보고서 한 줄 형식으로 바꿔 주는 **흐름 밖 화면**이다.
# 그래서 지키는 선이 셋이다:
#
#  1. **아무것도 저장하지 않는다.** 회차 파일·수시 카드·라벨·숨김·배정(group_overrides)·
#     소제목 순서(group_order)·커스텀 소제목 — 하나도 건드리지 않는다. 담당자가 들고 온
#     표는 이 앱이 모은 기사가 아니라서, 저장하는 순간 「앱이 모은 것」과 섞여 어느 쪽이
#     관측 기록인지 알 수 없게 된다. 받아서 바꿔 주고 끝낸다.
#  2. **텍스트 모양은 정기 값을 그대로 쓴다**(CLAUDE.md 「텍스트 형식」의 흐름 밖 규칙) —
#     기사 줄은 `article_line_template`, 소제목은 `subheading_format_template`. 이 화면
#     전용 형식 설정을 따로 만들지 않는다.
#  3. **AI는 담당자가 버튼을 눌렀을 때만** 부른다(자동 호출 0회). 정기 회차의 이름 이어쓰기
#     힌트와 섞이지 않게 llm_classifier.classify_isolated를 거친다.
from __future__ import annotations

import datetime as _dt
import io
import logging

from openpyxl import load_workbook

from app.config import MAX_SUBHEADINGS
from app.llm_classifier import (
    MAX_ARTICLES_FOR_CLASSIFY,
    UNCLASSIFIED_GROUP_NAME,
    assign_to_existing,
    classify_isolated,
    is_configured as llm_is_configured,
    last_assign_failure_reason,
    last_classification_failed_credit_error,
    last_classification_skipped_too_many_articles,
)

logger = logging.getLogger(__name__)

# 받는 표의 상한. 자르는 값이 아니라 **거절하는 값**이다 — 조용히 잘라 주면 담당자는
# 목록이 다 담긴 줄 알고 보고서를 보낸다(기관 기사 찾기의 EAGER_ROWS가 「미루는 값」인
# 것과 반대다. 여긴 미룰 뒷부분이 없다).
MAX_TABLE_ROWS = 1000
MAX_TABLE_COLS = 20
# 올린 .xlsx 원본 크기 상한 — 1000줄짜리 표는 보통 50KB 안쪽이다.
MAX_UPLOAD_BYTES = 4 * 1024 * 1024


class TableReadError(Exception):
    """표를 읽지 못한 이유 — 화면에 그대로 적는다(담당자가 할 일을 알 수 있게)."""


def _cell_text(value) -> str:
    """셀 하나를 글자로. 날짜·시각은 엑셀이 들고 있는 그대로가 아니라 사람이 적는 모양으로.

    openpyxl은 날짜 칸을 datetime으로 돌려주는데, `str()`을 그냥 쓰면 발행일이
    `2026-10-21 00:00:00`이 된다 — 시각이 0시인 건 「0시 발행」이 아니라 「시각이 없는
    날짜 칸」이라서, 그 경우엔 날짜만 적는다.
    """
    if value is None:
        return ""
    if isinstance(value, _dt.datetime):
        if value.hour or value.minute:
            return value.strftime("%Y-%m-%d %H:%M")
        return value.strftime("%Y-%m-%d")
    if isinstance(value, _dt.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_xlsx_rows(raw: bytes) -> list[list[str]]:
    """올린 .xlsx의 **첫 시트**를 2차원 글자 표로 읽는다.

    `data_only=True` — 수식 칸은 엑셀이 저장해 둔 계산 결과를 읽는다(수식 자체를 보고서에
    적을 수는 없다). 결과가 저장돼 있지 않은 파일이면 그 칸은 빈 칸이 된다.

    **셀에 걸린 하이퍼링크를 맨 뒤 열 하나로 꺼내 준다.** 엑셀에서 받은 표는 링크가 제목
    칸에 걸려 있고 주소 열이 따로 없는 경우가 흔한데, 그대로 읽으면 주소가 통째로 사라진다
    (이 화면은 링크가 없는 기사를 건너뛰므로 「표를 올렸는데 0건」이 된다). 열을 더해 두면
    담당자가 ② 열 맞추기에서 그 열을 「링크」로 고르거나 안 쓰면 그만이다 — 앱이 몰래
    고쳐 넣는 게 아니라 **보이는 열**로 내놓는 쪽을 쓴다.
    """
    if not raw:
        raise TableReadError("파일이 비어 있어요.")
    if len(raw) > MAX_UPLOAD_BYTES:
        raise TableReadError(f"파일이 너무 커요({len(raw) // 1024}KB). 표만 남겨 다시 저장해 주세요.")
    try:
        wb = load_workbook(io.BytesIO(raw), data_only=True)
    except Exception as exc:  # zipfile·xml 오류 등 — openpyxl이 종류를 가려주지 않는다
        logger.warning("엑셀 읽기 실패: %s", exc)
        raise TableReadError("엑셀 파일을 읽지 못했어요. .xlsx로 저장한 파일인지 확인해 주세요.")

    ws = wb[wb.sheetnames[0]]
    rows: list[list[str]] = []
    links: list[str] = []
    for row in ws.iter_rows():
        cells = [_cell_text(c.value) for c in row[:MAX_TABLE_COLS]]
        link = ""
        for c in row[:MAX_TABLE_COLS]:
            target = getattr(c.hyperlink, "target", None) if c.hyperlink else None
            if target and str(target).startswith("http") and not _cell_text(c.value).startswith("http"):
                link = str(target)
                break
        if not any(cells) and not link:
            continue  # 가운데 빈 줄은 건너뛴다(엑셀 표엔 흔하다)
        rows.append(cells)
        links.append(link)
        if len(rows) > MAX_TABLE_ROWS:
            raise TableReadError(
                f"표가 {MAX_TABLE_ROWS}줄을 넘어요. 나눠서 올려 주세요 — 잘라서 보여 주면 담당자가 다 담긴 줄 압니다."
            )
    wb.close()

    if not rows:
        raise TableReadError("표에 내용이 없어요.")
    if any(links):
        width = max(len(r) for r in rows)
        for r, link in zip(rows, links):
            r.extend([""] * (width - len(r)))
            r.append(link)
    return rows


# 「AI 소제목 배정」이 돌려주는 실패 이유 → 화면이 그대로 적을 한 줄. 모르는 이유는
# 일반 문구로 떨어진다(조용히 아무 일도 안 일어난 것처럼 보이지 않게).
ASSIGN_REASONS = {
    "no_key": "AI 키가 등록돼 있지 않아요. 설정 → LLM(AI) 연동에서 키를 넣어 주세요.",
    "nothing": "소제목이 비어 있는 기사가 없어요.",
    "too_many": f"기사가 {MAX_ARTICLES_FOR_CLASSIFY}건을 넘으면 AI가 한 번에 못 봐요. 몇 건 줄여 주세요.",
    "credit": "AI 크레딧이 부족해요. 충전 전에는 다시 눌러도 같은 결과입니다.",
    "no_match": "지금 있는 소제목 중에 맞는 곳이 없다고 답했어요. 소제목을 비우고 처음부터 배정해 보세요.",
    "failed": "AI를 부르지 못했어요. 잠시 뒤 다시 눌러 주세요.",
}


def assign_subheadings(items: list[dict], max_subheadings: int = MAX_SUBHEADINGS) -> dict:
    """표의 기사에 소제목을 붙인다. {"assigned": {url: 이름}} 또는 {"reason": 키}.

    **이미 소제목이 적힌 기사는 한 글자도 안 건드린다.** 그래서 하는 일이 둘로 갈린다 —
    정기 초안의 「AI 기사 배정」·수시 카드의 「소제목 배정」과 **같은 안전장치**다
    (CODING_CONVENTIONS.md §6-3 「재분류는 파괴적이라서 아낀다」):

      - 소제목이 하나도 없으면(표에 소제목 열이 없거나 전부 빈 칸) 처음부터 묶는다.
      - 하나라도 있으면 **빈 칸만** 그 이름들에 끼워 넣는다(assign_to_existing).

    모델이 응답에서 빠뜨린 기사는 전부 「기타」로 보낸다 — 한 번 눌러 미분류가 항상
    비도록(정기·수시의 같은 처리). 남은 자리를 넘는 새 이름도 「기타」로 간다.

    items: {"url", "title", "outlet", "summary", "section"} — summary는 표에 요약 열을
    맞춰 뒀을 때만 차 있다(없으면 제목만으로 묶는다. 요약이 있으면 훨씬 정확하다).
    """
    if not llm_is_configured():
        return {"reason": "no_key"}

    articles = [a for a in items if (a.get("url") or "").strip() and (a.get("title") or "").strip()]
    if not articles:
        return {"reason": "nothing"}
    if len(articles) > MAX_ARTICLES_FOR_CLASSIFY:
        return {"reason": "too_many"}

    unclassified = [a for a in articles if not (a.get("section") or "").strip()]
    if not unclassified:
        return {"reason": "nothing"}

    # 후보는 「표에 이미 적힌 소제목」. 「기타」는 배정 후보가 아니다(늘 열려 있는 칸이라
    # assign_to_existing이 애초에 받지 않는다).
    existing: list[str] = []
    for a in articles:
        name = (a.get("section") or "").strip()
        if name and name != "기타" and name not in existing:
            existing.append(name)
    room = max(0, max_subheadings - len(existing))

    if existing:
        assigned = assign_to_existing(unclassified, existing)
        if not assigned:
            reason = last_assign_failure_reason()
            return {"reason": "no_match" if reason == "no_match" else "failed"}
        accepted_new: set = set()
        out: dict = {}
        for a in unclassified:
            name = assigned.get(a["url"])
            if not name:
                out[a["url"]] = "기타"  # 응답에서 빠진 기사 — 미분류로 남기지 않는다
                continue
            if name not in existing and name not in accepted_new:
                if len(accepted_new) >= room:
                    name = "기타"
                else:
                    accepted_new.add(name)
            out[a["url"]] = name
        return {"assigned": out}

    groups = classify_isolated(articles, max_subheadings)
    if groups is None:
        if last_classification_skipped_too_many_articles():
            return {"reason": "too_many"}
        if last_classification_failed_credit_error():
            return {"reason": "credit"}
        return {"reason": "failed"}
    out = {}
    for g in groups:
        if g["name"] == UNCLASSIFIED_GROUP_NAME:
            continue  # 「📂 소제목 미분류」는 화면 전용 이름이라 보고서로 새면 안 된다
        for a in g["articles"]:
            out[a["url"]] = g["name"]
    for a in articles:
        out.setdefault(a["url"], "기타")
    return {"assigned": out}
