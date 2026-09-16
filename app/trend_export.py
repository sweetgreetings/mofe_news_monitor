# Design Ref: HISTORY.md "홈 화면 재구성" — 정책 단어 추이 표를 엑셀로 뽑는다.
"""정책 단어 추이(app.trend_renderer)의 단어×기간 표를 .xlsx로 만든다.

app.excel_export는 기사 8열(PRD.md 기능11) 고정 스키마를 5개 화면이 공유하는
모듈이라 여기서는 안 건드린다 — 이 표는 기사 목록이 아니라 완전히 다른 모양의
데이터(단어별 기간 집계)라 전용 빌더를 따로 둔다.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import COLOR_HEADER

_HEADER_FILL = PatternFill(start_color=COLOR_HEADER.lstrip("#"), end_color=COLOR_HEADER.lstrip("#"), fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def build_trend_workbook_bytes(bucket_labels: list, rows: list) -> bytes:
    """rows: [{"word": str, "values": [int|None, ...], "total": int}, ...].

    값이 None인 칸("기록 없음" — 그 구간에 저장된 회차가 아예 없음)은 0이 아니라
    빈 셀로 남긴다 — 0건과 다르다는 화면의 구분을 엑셀에서도 그대로 지킨다.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "정책 단어 추이"

    header = ["단어"] + list(bucket_labels) + ["합계"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL

    for row in rows:
        ws.append([row["word"]] + ["" if v is None else v for v in row["values"]] + [row["total"]])

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["A"].width = 18
    for i in range(2, 2 + len(bucket_labels) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 11

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
